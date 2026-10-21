import os
import io
import html
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional

try:
    from app.config import settings
except ImportError:
    try:
        from backend.app.config import settings
    except ImportError:
        settings = None

logger = logging.getLogger("report_service")


def _clean_text(val: Any) -> str:
    """Safely escape and clean strings for ReportLab Paragraphs."""
    if val is None:
        return ""
    text = str(val).strip()
    # Escape XML entities for ReportLab XML parser
    return html.escape(text)


class ReportService:
    def generate_student_pdf_bytes(
        self,
        report_data: Dict[str, Any],
        student_name: str,
        assignment_type: str = "text",
        file_name: str = "",
        assignment_title: str = "",
    ) -> bytes:
        """Create a printable GradeWise-branded report, including real sandbox logs.

        All report content is escaped and normalized for the built-in PDF fonts
        so Unicode bullets and unsupported glyphs cannot turn into black squares.
        """
        import re
        import textwrap
        import unicodedata
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
            HRFlowable, Preformatted, KeepTogether,
        )

        ink = colors.HexColor("#202020")
        deep = colors.HexColor("#151515")
        muted = colors.HexColor("#676660")
        accent = colors.HexColor("#F2E58F")
        soft = colors.HexColor("#F7F6F0")
        rule = colors.HexColor("#D9D7CE")
        danger = colors.HexColor("#A13D37")
        success = colors.HexColor("#356B47")

        def normalize(value: Any) -> str:
            raw = str(value if value is not None else "")
            for original, replacement in {
                "\u2022": "-", "\u25cf": "-", "\u25aa": "-", "\u25cf": "-",
                "\u2014": "-", "\u2013": "-", "\u2212": "-", "\u2192": "->",
                "\u2713": "PASS", "\u2717": "FAIL", "\u2714": "PASS",
                "\u201c": '"', "\u201d": '"', "\u2018": "'", "\u2019": "'",
                "\u00a0": " ", "\u200b": "", "\u2026": "...",
            }.items():
                raw = raw.replace(original, replacement)
            raw = unicodedata.normalize("NFKD", raw)
            # Helvetica's supported characters are essentially WinAnsi; emit
            # ASCII-only text for predictable PDF rendering across viewers.
            return raw.encode("ascii", "replace").decode("ascii").replace("\ufffd", "?")

        def para(value: Any, style, max_chars: int = 15000):
            text = normalize(value)[:max_chars]
            escaped = html.escape(text)
            return Paragraph(escaped.replace("\n", "<br/>"), style)

        styles = {
            "kicker": ParagraphStyle(
                "Gwkicker", fontName="Helvetica-Bold", fontSize=8,
                textColor=muted, leading=11, spaceAfter=4,
            ),
            "heading": ParagraphStyle(
                "Gwheading", fontName="Helvetica-Bold", fontSize=11,
                textColor=deep, leading=14, spaceBefore=13, spaceAfter=5,
            ),
            "body": ParagraphStyle(
                "Gwbody", fontName="Helvetica", fontSize=8.8,
                textColor=ink, leading=13.5, spaceAfter=5,
                splitLongWords=1, wordWrap="CJK",
            ),
            "small": ParagraphStyle(
                "Gwsmall", fontName="Helvetica", fontSize=7.7,
                textColor=ink, leading=11, splitLongWords=1,
                wordWrap="CJK",
            ),
            "label": ParagraphStyle(
                "Gwlabel", fontName="Helvetica-Bold", fontSize=8,
                textColor=muted, leading=11,
            ),
            "score": ParagraphStyle(
                "Gwscore", fontName="Helvetica-Bold", fontSize=23,
                textColor=deep, leading=27,
            ),
            "mono": ParagraphStyle(
                "Gwmono", fontName="Courier", fontSize=7.2,
                textColor=ink, leading=10, splitLongWords=1,
            ),
        }
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=letter,
            rightMargin=40, leftMargin=40, topMargin=95, bottomMargin=46,
            title="GradeWise | Student Feedback",
            author="GradeWise",
        )
        story: list = []

        def decorate(canvas, pdf_doc):
            canvas.saveState()
            width, height = letter
            canvas.setFillColor(deep)
            canvas.rect(0, height - 76, width, 76, fill=1, stroke=0)
            canvas.setFillColor(accent)
            canvas.rect(0, height - 80, width, 4, fill=1, stroke=0)
            canvas.setFont("Helvetica-Bold", 19)
            canvas.setFillColor(colors.HexColor("#F5F5F1"))
            canvas.drawString(41, height - 38, "GRADEWISE")
            canvas.setFillColor(accent)
            canvas.setFont("Helvetica-Bold", 8)
            canvas.drawString(42, height - 57, "AI-ASSISTED RUBRIC ASSESSMENT")
            canvas.setStrokeColor(rule)
            canvas.line(40, 33, width - 40, 33)
            canvas.setFillColor(muted)
            canvas.setFont("Helvetica", 8)
            canvas.drawString(40, 21, "GradeWise  /  Evidence-led feedback")
            canvas.drawRightString(width - 40, 21, f"Page {pdf_doc.page}")
            canvas.restoreState()

        score = float(report_data.get("score") if report_data.get("score") is not None
                      else report_data.get("overallScore") or 0)
        maximum = float(report_data.get("max_score") if report_data.get("max_score") is not None
                        else report_data.get("maxScore") or 100)
        percent = score / maximum * 100 if maximum else 0
        verdict = str(report_data.get("verdict") or report_data.get("status") or "Needs review")
        verdict = verdict.replace("_", " ").title()

        story.append(para("STUDENT FEEDBACK REPORT", styles["kicker"]))
        similarity = report_data.get("plagiarism_similarity_score", report_data.get("plagiarismScore"))
        integrity = (
            "Not checked (coming soon)" if settings and settings.DOCUMENT_DEMO
            else "Not checked"
        ) if similarity is None else f"{float(similarity):.1f}% Similarity"
        meta = [
            [para("Student", styles["label"]), para(student_name or "Student", styles["small"])],
            [para("Assignment", styles["label"]), para(assignment_title or report_data.get("assignmentTitle") or "Assignment", styles["small"])],
            [para("Submission", styles["label"]), para(file_name or report_data.get("fileName") or "Submission", styles["small"])],
            [para("Format", styles["label"]), para((assignment_type or "text").replace("_", " ").title(), styles["small"])],
            [para("Integrity", styles["label"]), para(integrity, styles["small"])],
            [para("Generated", styles["label"]), para(datetime.now().strftime("%d %b %Y, %H:%M"), styles["small"])],
        ]
        meta_table = Table(meta, colWidths=[95, 437], hAlign="LEFT")
        meta_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), soft),
            ("LINEBELOW", (0, -1), (-1, -1), 0.5, rule),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 11))
        scorebox = Table([
            [para(f"{score:.1f} / {maximum:.1f}", styles["score"]),
             para(f"{percent:.1f}%   |   {verdict}", styles["heading"])],
        ], colWidths=[210, 322], hAlign="LEFT")
        scorebox.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), accent),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 12),
            ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ("TOPPADDING", (0, 0), (-1, -1), 10),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(scorebox)

        summary = (
            report_data.get("student_feedback") or report_data.get("summary")
            or report_data.get("feedback") or report_data.get("instructor_brief")
        )
        if summary:
            story.append(para("SUMMARY & FEEDBACK", styles["heading"]))
            story.append(para(summary, styles["body"]))

        scores = report_data.get("per_criterion") or {}
        notes = report_data.get("criteria_notes") or {}
        if scores:
            story.append(para("RUBRIC EVALUATION", styles["heading"]))
            for criterion_id, raw_score in scores.items():
                info = notes.get(criterion_id) or {}
                if isinstance(raw_score, dict):
                    info = {**raw_score, **info}
                    raw_score = raw_score.get("score", 0)
                try:
                    pts = f"{float(raw_score):.1f} pts"
                except (ValueError, TypeError):
                    pts = "Reviewed"
                heading = Table([
                    [para(str(criterion_id).replace("_", " ").title(), styles["label"]),
                     para(pts, styles["label"])],
                ], colWidths=[400, 132], hAlign="LEFT")
                heading.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), soft),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.8, accent),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(Spacer(1, 5))
                story.append(heading)
                for field, label in [
                    ("evidence", "Evidence"), ("deduction_reason", "Deductions"),
                    ("reference_support", "Reference support"), ("action", "Recommended action"),
                    ("note", "Assessment"),
                ]:
                    if info.get(field):
                        story.append(para(f"{label}: {info[field]}", styles["small"]))

        for key, title in [
            ("strengths", "STRENGTHS"),
            ("weaknesses", "AREAS FOR IMPROVEMENT"),
            ("priority_actions", "PRIORITY ACTIONS"),
            ("suggestions", "RECOMMENDATIONS"),
        ]:
            items = report_data.get(key) or []
            if not isinstance(items, list) or not items:
                continue
            story.append(para(title, styles["heading"]))
            for item in items:
                if str(item).strip():
                    # ASCII hyphen instead of the Helvetica-incompatible Unicode bullet.
                    story.append(para("- " + str(item), styles["body"]))

        def log_block(label, content):
            story.append(para(label, styles["label"]))
            text = normalize(content)
            if not text.strip():
                text = "(empty)"
            lines = []
            for line in text.splitlines() or [text]:
                lines.extend(textwrap.wrap(
                    line, width=92, replace_whitespace=False,
                    drop_whitespace=False, break_long_words=True,
                    break_on_hyphens=False,
                ) or [""])
            # Split the log in small blocks so ReportLab can paginate safely.
            for offset in range(0, len(lines), 28):
                story.append(Preformatted(
                    "\n".join(lines[offset:offset + 28]), styles["mono"],
                    maxLineLength=92, splitChars=" ",
                ))
                story.append(Spacer(1, 3))

        structured_logs = report_data.get("sandbox_logs")
        legacy_log = report_data.get("code_execution_result")
        if isinstance(structured_logs, list) and structured_logs:
            story.append(para("SANDBOX EXECUTION LOGS", styles["heading"]))
            story.append(para(
                "Exit code 0 means the program ran; correctness is assessed against the rubric.",
                styles["small"],
            ))
            for index, item in enumerate(structured_logs[:12], 1):
                if not isinstance(item, dict):
                    continue
                story.append(para(
                    f"TEST {index}: {item.get('name') or 'Execution'} - "
                    f"exit {item.get('exitCode', 'unknown')} - "
                    f"{item.get('durationMs', 'unknown')} ms",
                    styles["heading"],
                ))
                log_block("Input (stdin)", item.get("input", ""))
                log_block("Output (stdout)", item.get("stdout", item.get("actual_output", "")))
                log_block("Errors (stderr)", item.get("stderr", item.get("error", "")))
        elif isinstance(structured_logs, str) and structured_logs.strip():
            story.append(para("SANDBOX EXECUTION LOGS", styles["heading"]))
            log_block("Recorded sandbox input / output / errors", structured_logs)
        elif isinstance(legacy_log, str) and legacy_log.strip():
            story.append(para("SANDBOX EXECUTION LOGS", styles["heading"]))
            log_block("Recorded sandbox input / output / errors", legacy_log)
        elif assignment_type == "coding":
            story.append(para("SANDBOX EXECUTION LOGS", styles["heading"]))
            story.append(para("No execution logs were stored for this submission.", styles["small"]))

        conclusion = report_data.get("score_rationale") or report_data.get("finalComments")
        if conclusion and conclusion != summary:
            story.append(para("EVALUATOR COMMENTS", styles["heading"]))
            story.append(para(conclusion, styles["body"]))

        doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
        result = buffer.getvalue()
        buffer.close()
        return result

    async def generate_student_pdf(
        self,
        report_data: Dict[str, Any],
        student_name: str,
        assignment_type: str,
        output_path: str,
        file_name: str = "",
        assignment_title: str = "",
    ) -> str:
        """
        Generates PDF report and writes to output_path.
        """
        logger.info(f"Generating student PDF report file for {student_name} at {output_path}...")
        try:
            pdf_bytes = self.generate_student_pdf_bytes(
                report_data=report_data,
                student_name=student_name,
                assignment_type=assignment_type,
                file_name=file_name,
                assignment_title=assignment_title,
            )
            with open(output_path, "wb") as f:
                f.write(pdf_bytes)
            logger.info("Successfully generated PDF report via ReportLab.")
            return output_path
        except Exception as e:
            logger.error(f"ReportLab PDF compilation failed: {e}. Writing markdown file fallback instead.")
            md_path = output_path.replace(".pdf", ".md")
            self._write_markdown_report(report_data, student_name, assignment_type, md_path)
            return md_path

    async def generate_batch_excel(self, batch_data: List[Dict[str, Any]], output_path: str) -> str:
        """
        Compiles batch job grades into an Excel worksheet.
        """
        logger.info(f"Generating batch Excel overview at {output_path}...")
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Batch Grades Report"

            headers = [
                "Student Name", "File Name", "Score", "Max Score", "Percentage",
                "Verdict / Status", "Strengths Summary", "Weaknesses Summary",
                "Actionable Suggestions", "Plagiarism %"
            ]
            ws.append(headers)

            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center

            for row_idx, record in enumerate(batch_data, start=2):
                strengths = "; ".join(record.get("strengths", [])) if isinstance(record.get("strengths"), list) else str(record.get("strengths", ""))
                weaknesses = "; ".join(record.get("weaknesses", [])) if isinstance(record.get("weaknesses"), list) else str(record.get("weaknesses", ""))
                suggestions = "; ".join(record.get("suggestions", [])) if isinstance(record.get("suggestions"), list) else str(record.get("suggestions", ""))

                score = float(record.get("score") if record.get("score") is not None else record.get("overallScore", 0))
                max_score = float(record.get("max_score") if record.get("max_score") is not None else record.get("maxScore", 100))
                pct = f"{(score / max_score * 100):.1f}%" if max_score > 0 else "0%"

                row_data = [
                    record.get("studentName", record.get("student_name", "Student")),
                    record.get("fileName", record.get("file_name", "assignment")),
                    score,
                    max_score,
                    pct,
                    str(record.get("status", record.get("verdict", "COMPLETED"))).upper(),
                    strengths,
                    weaknesses,
                    suggestions,
                    f"{record.get('plagiarismScore', record.get('plagiarism_similarity_score', 0.0))}%",
                ]
                ws.append(row_data)

                for col_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.alignment = align_center if col_idx in [3, 4, 5, 6, 10] else align_left

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

            wb.save(output_path)
            logger.info("Successfully generated Excel report via openpyxl.")
            return output_path

        except Exception as e:
            logger.error(f"Excel compilation failed: {e}. Writing standard CSV sheet instead.")
            csv_path = output_path.replace(".xlsx", ".csv")
            self._write_csv_report(batch_data, csv_path)
            return csv_path

    def _write_markdown_report(self, report_data: Dict[str, Any], student_name: str, assignment_type: str, path: str):
        content = (
            f"# AI-Powered Feedback Report\n\n"
            f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            f"**Student Name:** {student_name}\n"
            f"**Assignment Type:** {assignment_type.capitalize()}\n"
            f"**Overall Grade:** {report_data.get('score', report_data.get('overallScore', 0))}/{report_data.get('max_score', report_data.get('maxScore', 100))}\n"
            f"**Status:** {report_data.get('status', 'Completed').upper()}\n\n"
            f"## Performance Summary\n"
            f"{report_data.get('summary', report_data.get('student_feedback', ''))}\n\n"
            f"## Key Strengths\n"
        )
        for s in report_data.get("strengths", []):
            content += f"- {s}\n"

        content += "\n## Areas for Improvement\n"
        for w in report_data.get("weaknesses", []):
            content += f"- {w}\n"

        content += "\n## Actionable Suggestions\n"
        for sug in report_data.get("suggestions", []):
            content += f"- {sug}\n"

        with open(path, "w", encoding="utf-8") as f:
            f.write(content)

    def _write_csv_report(self, batch_data: List[Dict[str, Any]], path: str):
        import csv
        headers = [
            "Student Name", "File Name", "Overall Score", "Status",
            "Strengths Summary", "Weaknesses Summary", "Suggestions Summary",
            "Plagiarism %"
        ]
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for record in batch_data:
                writer.writerow([
                    record.get("studentName", record.get("student_name", "Student")),
                    record.get("fileName", record.get("file_name", "assignment")),
                    record.get("score", record.get("overallScore", 0)),
                    record.get("status", "completed").upper(),
                    "; ".join(record.get("strengths", [])) if isinstance(record.get("strengths"), list) else str(record.get("strengths", "")),
                    "; ".join(record.get("weaknesses", [])) if isinstance(record.get("weaknesses"), list) else str(record.get("weaknesses", "")),
                    "; ".join(record.get("suggestions", [])) if isinstance(record.get("suggestions"), list) else str(record.get("suggestions", "")),
                    f"{record.get('plagiarismScore', 0.0)}%",
                ])


# Initialize global instance
report_service = ReportService()
